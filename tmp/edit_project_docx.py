from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(
    r"C:\Users\caleb\Documents\Estudios\Universidad\TECNICO\CICLO02\DSP\DSP Teoria\Documentación de Proyecto de Cátedra.docx"
)
OUTPUT = Path(
    r"C:\Users\caleb\Documents\Estudios\Universidad\TECNICO\CICLO02\DSP\DSP Teoria\DSP Proyecto Catedra\output\documents\Documentación de Proyecto de Cátedra - corregida.docx"
)


def replace_run_text(paragraph, old, new):
    replaced = False
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            replaced = True
    if not replaced:
        raise ValueError(f"Text not found in paragraph: {old!r}")


def set_paragraph_text_preserving_first_run(paragraph, text, size_pt=None):
    if not paragraph.runs:
        run = paragraph.add_run(text)
    else:
        run = paragraph.runs[0]
        run.text = text
        for extra in paragraph.runs[1:]:
            extra.text = ""
    if size_pt is not None:
        run.font.size = Pt(size_pt)


OUTPUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(SOURCE, OUTPUT)
doc = Document(OUTPUT)

# Cover: add the actual project title and make each member-role assignment explicit.
set_paragraph_text_preserving_first_run(
    doc.paragraphs[2],
    "Sistema web inteligente para la gestión de renta de vehículos - Fase 1",
    12,
)

member_lines = {
    6: "Flores Figueroa, Josue Gamaliel (FF233029) - Líder de proyecto y analista funcional",
    7: "Torres Reyes, Cristian Josué (TR240516) - Desarrollador Backend e IA",
    8: "Ortiz Melendez, Michael Caleb (OM260275) - Desarrollador Frontend y UX/UI",
    9: "Martinez Guerrero, Nathaly Ivania (MG260121) - Base de datos y DevOps",
    10: "Peña Saravia, Jimmy Steeven (PS260123) - Pruebas, calidad y documentación",
}
for index, text in member_lines.items():
    paragraph = doc.paragraphs[index]
    set_paragraph_text_preserving_first_run(paragraph, text, 10.5)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
doc.paragraphs[10].paragraph_format.space_after = Pt(6)

replace_run_text(doc.paragraphs[16], "24 de Agosto de 2026", "24 de agosto de 2026")

# Minimal language corrections found during the rubric review.
corrections = {
    27: ("e incorpora un asistente", "e incorporar un asistente"),
    163: ("delegaron los casos de uso", "delegarán los casos de uso"),
    180: (
        "Se registran versión del prompt, candidatos y resultado para pruebas.",
        "Se registrarán la versión del prompt, los candidatos y el resultado para las pruebas.",
    ),
    182: ("La aplicación se ampliará en una imagen Docker", "La aplicación se empaquetará como una imagen Docker"),
    185: ("El servicio administrador de PostgreSQL", "El servicio administrado de PostgreSQL"),
    269: ("Los hallazgos se organizan", "Los hallazgos se organizarán"),
    290: ("RNF-07. Ejecuta de manera reproducible", "RNF-07. Ejecutarse de manera reproducible"),
    293: ("Los requisitos se priorizará", "Los requisitos se priorizarán"),
    295: ("Primero se valorarán los flujos", "Primero se validarán los flujos"),
    333: ("Semana 4: Implementar", "Semana 4: implementar"),
    336: ("Semana 7: Integrar", "Semana 7: integrar"),
    379: (
        "Microsoft, “Entity Framework documentation hub\", \" Microsoft Learn.",
        "Microsoft, “Entity Framework documentation hub,” Microsoft Learn.",
    ),
}
for index, (old, new) in corrections.items():
    replace_run_text(doc.paragraphs[index], old, new)

# Layout repairs found in the rendered review.
# Keep the cover separate from the table of contents.
doc.paragraphs[18].paragraph_format.page_break_before = True

# This image paragraph had accidentally inherited Heading 3, causing Word to
# copy the use-case diagram into the table of contents.
doc.paragraphs[151].style = doc.styles["Normal"]
doc.paragraphs[151].alignment = 1

# Keep short diagram labels with the figures that follow them.
doc.paragraphs[149].paragraph_format.keep_with_next = True
doc.paragraphs[150].paragraph_format.keep_with_next = True
doc.paragraphs[152].paragraph_format.keep_with_next = True
doc.paragraphs[155].paragraph_format.keep_with_next = True
doc.paragraphs[157].paragraph_format.keep_with_next = True

# The activity diagram exceeded the printable height, leaving its label alone
# on one page. Scale it proportionally so label and figure stay together.
activity_diagram = doc.inline_shapes[7]
activity_diagram.height = Inches(9.0)
activity_diagram.width = Inches(5.14)

# Repeat column labels when long tables continue onto another page.
for table in doc.tables:
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:tblHeader"))
    if existing is None:
        existing = OxmlElement("w:tblHeader")
        tr_pr.append(existing)
    existing.set(qn("w:val"), "true")

# Ask Word-compatible editors to refresh the TOC/page-number fields when opened.
settings = doc.settings.element
for existing in settings.findall(qn("w:updateFields")):
    settings.remove(existing)
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
settings.append(update_fields)

doc.save(OUTPUT)
print(OUTPUT)
